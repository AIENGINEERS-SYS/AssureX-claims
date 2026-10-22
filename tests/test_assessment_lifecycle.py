"""Integrated submission, durable workers, reports and adversarial access checks."""
from datetime import timedelta
from io import BytesIO
import json
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import func, select
from backend.db.assessment_models import AssessmentAttempt, AssessmentJob, AssessmentReport
from backend.db.models import Claim, EvaluationResult, Notification, NotificationPreference, utcnow
from backend.extensions import db
from backend.services import assessments, assessment_reports
from test_auth import app, client, accounts, claims, login, bearer
from test_claim_submission import complete, submit


@pytest.fixture(autouse=True)
def model_configuration(app, tmp_path):
    app.config.update(GTM_PREDICTOR=lambda path: [.85, .1, .05],
        MODEL_CARD_PATH=str(tmp_path / "cards"), ASSESSMENT_MAX_ATTEMPTS=2,
        ASSESSMENT_RETRY_SECONDS=1)


def auth(client, role="customer"):
    return bearer(login(client, role)["access_token"])


def schedule(app, claim_id):
    with app.app_context():
        job = assessments.enqueue(db.session.get(Claim, claim_id))
        db.session.commit()
        return job.id


def work(app):
    with app.app_context():
        assert assessments.process_next()


def ready_again(app, identity):
    with app.app_context():
        job = db.session.get(AssessmentJob, identity)
        job.available_at = utcnow() - timedelta(seconds=1)
        db.session.commit()


def test_submission_runs_in_worker_and_customer_downloads_saved_report(app, client, accounts, claims):
    headers = auth(client)
    draft = complete(client, headers, claims["customer"]["product"])
    response = submit(client, headers, draft)
    assert response.status_code == 201, response.json
    claim = response.json["claim"]
    assert claim["assessment"]["status"] == "queued"
    with app.app_context():
        assert db.session.scalar(select(func.count()).select_from(EvaluationResult)) == 0
    assert submit(client, headers, draft).json["claim"]["id"] == claim["id"]
    work(app)
    progress = client.get(f"/api/claims/{claim['id']}/assessment", headers=headers).json
    assert progress["assessment"]["assessment_status"] == "completed", progress
    assert progress["assessment"]["report_status"] == "available"
    report_id = progress["assessment"]["report"]["id"]
    url = f"/api/claims/{claim['id']}/assessment/reports/{report_id}"
    report = client.get(url, headers=headers).json
    payload = report["report"]["payload"]
    assert payload["claim"]["claim_id"] == claim["claim_id"]
    assert len(payload["evidence_inventory"]) == 4
    assert payload["assessment_version"] == 1 and payload["evaluation_id"]
    assert payload["authorized_decision"] is None
    assert "not a final" in payload["notice"] or "requires an authorized" in payload["notice"]
    for format in ("pdf", "csv", "excel"):
        artifact = client.get(url + "/download/" + format, headers=headers)
        assert artifact.status_code == 200, artifact.json
        if format == "pdf":
            assert claim["claim_id"] in "".join(page.extract_text() for page in PdfReader(BytesIO(artifact.data)).pages)
        elif format == "excel":
            workbook = load_workbook(BytesIO(artifact.data), read_only=True)
            assert "claim" in workbook.sheetnames
            workbook.close()
        else:
            assert claim["claim_id"].encode() in artifact.data
    with app.app_context():
        assert db.session.scalar(select(func.count()).select_from(AssessmentJob)) == 1
        assert db.session.scalar(select(func.count()).select_from(AssessmentReport)) == 1
        assert db.session.scalar(select(func.count()).select_from(AssessmentAttempt)) == 1
        assert db.session.get(Claim, claim["id"]).final_decision is None
        assert db.session.scalar(select(func.count()).select_from(Notification).where(Notification.type == "REPORT_AVAILABLE")) == 1


def test_report_failure_retries_without_repeating_inference(app, client, claims, monkeypatch):
    identity = schedule(app, claims["customer"]["claim"])
    original = assessment_reports.generate
    monkeypatch.setattr(assessment_reports, "generate", lambda *args: (_ for _ in ()).throw(RuntimeError("private traceback /secret")))
    work(app)
    with app.app_context():
        job = db.session.get(AssessmentJob, identity)
        assert job.status == "retry_wait" and job.report_status == "failed"
        assert job.evaluation_snapshot["status"] == "complete"
        evaluation_id = job.evaluation_id
        assert db.session.scalar(select(func.count()).select_from(EvaluationResult)) == 1
    public = client.get(f"/api/claims/{claims['customer']['claim']}/assessment", headers=auth(client)).json
    assert "secret" not in json.dumps(public)
    assert public["assessment"]["assessment_status"] == "completed"
    monkeypatch.setattr(assessment_reports, "generate", original)
    ready_again(app, identity)
    work(app)
    with app.app_context():
        job = db.session.get(AssessmentJob, identity)
        assert job.status == "completed" and job.evaluation_id == evaluation_id
        assert job.attempts == 1 and job.report_attempts == 2
        assert db.session.scalar(select(func.count()).select_from(EvaluationResult)) == 1


def test_partial_model_failure_has_bounded_retries_and_no_success_report(app, client, claims):
    app.config["GTM_PREDICTOR"] = lambda path: (_ for _ in ()).throw(RuntimeError("secret model directory"))
    claim_id = claims["customer"]["claim"]
    identity = schedule(app, claim_id)
    work(app)
    ready_again(app, identity)
    work(app)
    with app.app_context():
        job = db.session.get(AssessmentJob, identity)
        assert job.status == "failed" and job.attempts == 2
        assert job.evaluation_snapshot["status"] == "partial"
        assert db.session.scalar(select(func.count()).select_from(AssessmentReport)) == 0
        assert db.session.scalar(select(func.count()).select_from(AssessmentAttempt)) == 2
        assert db.session.get(Claim, claim_id).final_decision is None
    data = client.get(f"/api/claims/{claim_id}/assessment", headers=auth(client)).json
    assert data["assessment"]["retryable"] and "secret" not in json.dumps(data)


def test_worker_lease_recovery_fences_old_worker_and_prevents_double_claim(app, claims):
    identity = schedule(app, claims["customer"]["claim"])
    with app.app_context():
        old = assessments.take(identity)
        assert old and assessments.take(identity) is None
        job = db.session.get(AssessmentJob, identity)
        job.lease_until = utcnow() - timedelta(seconds=1)
        db.session.commit()
        assessments.recover()
    ready_again(app, identity)
    with app.app_context():
        new = assessments.take(identity)
        assert new and old[1] != new[1]
        assessments.run(*old)
        assert db.session.get(AssessmentJob, identity).status == "running"
        assessments.run(*new)
        assert db.session.scalar(select(func.count()).select_from(AssessmentReport)) == 1


def test_evidence_reassessment_preserves_historical_report_and_final_decisions(app, client, claims):
    claim_id = claims["customer"]["claim"]
    first = schedule(app, claim_id)
    assert schedule(app, claim_id) == first
    work(app)
    with app.app_context():
        original = db.session.scalar(select(AssessmentReport))
        original_payload = json.dumps(original.payload, sort_keys=True)
        original_id = original.id
        claim = db.session.get(Claim, claim_id)
        claim.fault_description = "Updated fault description and newly verified evidence"
        db.session.commit()
    second = schedule(app, claim_id)
    assert second != first
    work(app)
    with app.app_context():
        assert json.dumps(db.session.get(AssessmentReport, original_id).payload, sort_keys=True) == original_payload
        assert db.session.scalar(select(func.count()).select_from(AssessmentReport)) == 2
        claim = db.session.get(Claim, claim_id)
        claim.fault_description = json.loads(original_payload)["claim"]["fault_description"]
        db.session.commit()
    third = schedule(app, claim_id)
    assert third not in {first, second}
    work(app)
    with app.app_context():
        assert db.session.scalar(select(func.count()).select_from(AssessmentReport)) == 3
        assert json.dumps(db.session.get(AssessmentReport, original_id).payload, sort_keys=True) == original_payload
        claim = db.session.get(Claim, claim_id)
        claim.status, claim.manual_review_required = "manual_review", True
        db.session.commit()
    reviewer = auth(client, "reviewer")
    decision = client.post(f"/api/review/{claim_id}/override", headers=reviewer,
        json={"decision": "approve", "notes": "Verified original evidence", "override_reason": "Verified evidence directly with retailer"})
    assert decision.status_code == 200, decision.json
    assert client.post(f"/api/claims/{claim_id}/evaluate", headers=reviewer).status_code == 409
    assert client.post(f"/api/claims/{claim_id}/assessment/retry", headers=reviewer, json={}).status_code == 409
    latest = client.get(f"/api/claims/{claim_id}/assessment/reports/{original_id}", headers=auth(client)).json
    assert latest["authorized_decision"]["decision"] == "approve"
    assert latest["report"]["payload"]["authorized_decision"] is None


def test_access_checks_cover_nested_resources_reassignment_and_unassigned_review(app, client, claims):
    claim_id = claims["customer"]["claim"]
    schedule(app, claim_id)
    work(app)
    with app.app_context():
        report_id = db.session.scalar(select(AssessmentReport.id))
    paths = [f"/api/claims/{claim_id}/assessment", f"/api/claims/{claim_id}/assessment/reports",
             f"/api/claims/{claim_id}/assessment/reports/{report_id}",
             f"/api/claims/{claim_id}/assessment/reports/{report_id}/download/pdf"]
    other = auth(client, "other")
    employee = auth(client, "employee")
    reviewer = auth(client, "reviewer")
    for path in paths:
        assert client.get(path).status_code == 401
        assert client.get(path, headers=other).status_code == 404
        assert client.get(path, headers=employee).status_code == 200
    with app.app_context():
        claim = db.session.get(Claim, claim_id)
        claim.assigned_employee_id = claim.assigned_reviewer_id = None
        db.session.commit()
    for path in paths:
        assert client.get(path, headers=employee).status_code == 404
        assert client.get(path, headers=reviewer).status_code == 404
    assert client.post(f"/api/claims/{claim_id}/evaluate", headers=reviewer).status_code == 404
    assert client.patch(f"/api/dashboard/reviewer/claims/{claim_id}/assignment", headers=reviewer, json={}).status_code == 403


def test_report_notification_respects_preferences(app, accounts, claims):
    with app.app_context():
        db.session.add(NotificationPreference(user_id=accounts["customer"], claim_updates=False))
        db.session.commit()
    schedule(app, claims["customer"]["claim"])
    work(app)
    with app.app_context():
        assert db.session.scalar(select(func.count()).select_from(AssessmentReport)) == 1
        assert db.session.scalar(select(func.count()).select_from(Notification).where(Notification.type == "REPORT_AVAILABLE")) == 0


def test_evidence_changes_before_worker_are_explicitly_superseded(app, claims):
    identity = schedule(app, claims["customer"]["claim"])
    with app.app_context():
        db.session.get(Claim, claims["customer"]["claim"]).fault_type = "Changed evidence"
        db.session.commit()
    work(app)
    with app.app_context():
        job = db.session.get(AssessmentJob, identity)
        assert job.status == "superseded" and job.error_code == "evidence_changed"
        assert db.session.scalar(select(func.count()).select_from(EvaluationResult)) == 0


def test_registered_routes_have_no_duplicate_methods(app):
    seen = set()
    for rule in app.url_map.iter_rules():
        for method in rule.methods - {"HEAD", "OPTIONS"}:
            signature = (rule.rule, method)
            assert signature not in seen, signature
            seen.add(signature)


def test_requested_information_requires_new_requested_evidence_before_reassessment(app, client, accounts, claims):
    from io import BytesIO
    from test_claim_submission import png
    headers = auth(client)
    draft = complete(client, headers, claims["customer"]["product"])
    submitted = submit(client, headers, draft).json["claim"]
    work(app)
    claim_id = submitted["id"]
    admin = auth(client, "admin")
    assert client.patch(f"/api/dashboard/reviewer/claims/{claim_id}/assignment", headers=admin,
        json={"reviewer_id": accounts["reviewer"]}).status_code == 200
    reviewer = auth(client, "reviewer")
    requested = client.post(f"/api/dashboard/reviewer/claims/{claim_id}/request-documents", headers=reviewer,
        json={"document_types": ["diagnostic_report"]})
    assert requested.status_code == 200, requested.json
    details = client.get(f"/api/claims/{claim_id}", headers=headers).json["claim"]
    assert details["information_request"]["document_types"] == ["diagnostic_report"]
    assert "diagnostic_report" in details["document_policy"]["required_types"]
    assert "existing_document_ids" not in details["information_request"]
    missing = client.post(f"/api/claims/{claim_id}/assessment/retry", headers=headers, json={})
    assert missing.status_code == 400
    assert "diagnostic_report" in str(missing.json)
    uploaded = client.post(f"/api/claims/{claim_id}/documents", headers=headers, data={
        "version": str(details["version"]), "document_type": "diagnostic_report",
        "file": (BytesIO(png("purple")), "diagnostic.png", "image/png")})
    assert uploaded.status_code == 201, uploaded.json
    retried = client.post(f"/api/claims/{claim_id}/assessment/retry", headers=headers, json={})
    assert retried.status_code == 202, retried.json
    assert retried.json["claim"]["information_request"] == {}
    work(app)
    with app.app_context():
        assert db.session.scalar(select(func.count()).select_from(AssessmentReport)) == 2


def test_missing_python_artifact_keeps_gtm_result_and_is_not_success(app, claims, tmp_path):
    app.config["PYTHON_MODEL_PATH"] = str(tmp_path / "unavailable.joblib")
    identity = schedule(app, claims["customer"]["claim"])
    work(app)
    with app.app_context():
        job = db.session.get(AssessmentJob, identity)
        item = db.session.get(EvaluationResult, job.evaluation_id)
        assert item.status == "partial" and item.gtm_prediction_id is not None
        assert item.python_prediction_id is None
        assert job.status == "retry_wait" and job.report_status == "pending"


def test_assessment_migration_roundtrip_archives_legacy_machine_decision(app, claims):
    from alembic import command
    from alembic.config import Config
    from sqlalchemy import inspect
    from backend.db.models import AuditLog
    with app.app_context():
        claim_id = claims["customer"]["claim"]
        db.session.get(Claim, claim_id).final_decision = "likely_invalid"
        db.session.commit()
        db.session.remove()
        command.downgrade(Config("backend/db/alembic.ini"), "a35e91b70412")
        assert "assessment_jobs" not in inspect(db.engine).get_table_names()
        command.upgrade(Config("backend/db/alembic.ini"), "head")
        assert db.session.get(Claim, claim_id).final_decision is None
        event = db.session.scalar(select(AuditLog).where(AuditLog.action == "claim.automated_decision_reclassified"))
        assert event.old_values["final_decision"] == "likely_invalid"


def test_assessment_inputs_and_published_reports_cannot_be_rewritten(app, claims):
    identity = schedule(app, claims["customer"]["claim"])
    work(app)
    with app.app_context():
        job = db.session.get(AssessmentJob, identity)
        job.evidence_snapshot = {"changed": True}
        with pytest.raises(ValueError, match="immutable"):
            db.session.commit()
        db.session.rollback()
        report = db.session.scalar(select(AssessmentReport))
        report.payload = {"fabricated": "approval"}
        with pytest.raises(ValueError, match="append-only"):
            db.session.commit()
        db.session.rollback()
        assert db.session.get(AssessmentJob, identity).status == "completed"


def test_concurrent_scheduling_preserves_one_job(app, claims):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    from sqlalchemy.exc import OperationalError
    from werkzeug.exceptions import Conflict
    barrier = Barrier(2)
    claim_id = claims["customer"]["claim"]
    def enqueue_once():
        with app.app_context():
            claim = db.session.get(Claim, claim_id)
            barrier.wait(timeout=15)
            try:
                job = assessments.enqueue(claim)
                db.session.commit()
                return job.id
            except (Conflict, OperationalError):
                db.session.rollback()
                return None
    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(lambda _: enqueue_once(), range(2)))
    assert any(results)
    with app.app_context():
        assert db.session.scalar(select(func.count()).select_from(AssessmentJob)) == 1


def test_pending_evidence_cannot_be_reported_as_complete(app, claims, accounts):
    from fuzzy_helpers import add_document
    from backend.db.models import Document
    claim_id = claims["customer"]["claim"]
    public_id = add_document(app, claim_id, accounts["customer"], "receipt")
    with app.app_context():
        db.session.scalar(select(Document).where(Document.document_id == public_id)).ocr_status = "processing"
        db.session.commit()
    identity = schedule(app, claim_id)
    work(app)
    with app.app_context():
        job = db.session.get(AssessmentJob, identity)
        assert job.error_code == "evidence_processing" and job.status == "retry_wait"
        assert job.evaluation_id is None
        assert db.session.scalar(select(func.count()).select_from(AssessmentReport)) == 0
