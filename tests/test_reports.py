"""Report integration tests use real migrations, JWTs, workers and parsed artifacts."""
import csv
from datetime import date, timedelta
from io import BytesIO, StringIO
import json
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select, func
from backend.db.models import AuditLog, Claim, Review, User, Warranty, utcnow
from backend.db.report_models import ReportJob, ReportItem
from backend.extensions import db
from backend.services.report_jobs import process_next, cleanup
from test_auth import app, client, accounts, claims, login, bearer  # noqa: F401


@pytest.fixture(autouse=True)
def storage(app, tmp_path):
    app.config.update(REPORT_STORAGE_PATH=str(tmp_path / 'reports'), REPORT_BATCH_SIZE=1000)


def auth(client, role='customer'):
    return bearer(login(client, role)['access_token'])


def exported(app, client, headers, format='csv', filters=None):
    response = client.post(f'/api/reports/export/{format}', json=filters or {}, headers=headers)
    assert response.status_code == 202, response.json
    job_id = response.json['id']
    with app.app_context():
        assert process_next()
    status = client.get(f'/api/reports/jobs/{job_id}', headers=headers)
    assert status.json['status'] == 'completed', status.json
    assert status.json['progress'] == 100
    result = client.get(f'/api/reports/jobs/{job_id}/download', headers=headers)
    assert result.status_code == 200, result.json
    return job_id, result


def test_report_preview_permissions_and_ownership(app, client, accounts, claims):
    mine, other = claims['customer'], claims['other']
    assert client.get('/api/reports').status_code == 401
    for role in ['customer', 'employee', 'admin']:
        headers = auth(client, role)
        assert client.get(f'/api/reports/claim/{mine["claim"]}', headers=headers).status_code == 200
        assert client.get(f'/api/reports/claim/{other["claim"]}', headers=headers).status_code == (200 if role == 'admin' else 404)
    headers = auth(client)
    assert client.get(f'/api/reports/customer/{accounts["other"]}', headers=headers).status_code == 404
    assert client.get(f'/api/reports/product/{other["product"]}', headers=headers).status_code == 404
    assert client.get('/api/reports/administrative', headers=headers).status_code == 403
    assert client.post('/api/reports/export/csv', headers=headers, json={
        'report_type': 'claim', 'entity_id': other['claim']}).status_code == 404
    data = client.get(f'/api/reports/customer/{accounts["customer"]}', headers=headers).json
    assert data['sections']['Claims']['total'] == 1
    assert data['sections']['Warranties']['total'] == 1
    assert data['sections']['Products']['items'][0]['id'] == mine['product']
    assert 'password_hash' not in json.dumps(data) and 'storage_path' not in json.dumps(data)
    admin = client.get('/api/reports/administrative', headers=auth(client, 'admin')).json
    assert admin['analytics']['total_claims'] == 2


def test_customer_reports_hide_internal_review_and_model_details(app, client, accounts, claims):
    with app.app_context():
        claim = db.session.get(Claim, claims['customer']['claim'])
        db.session.add(Review(claim_id=claim.id, reviewer_user_id=accounts['reviewer'],
            decision='manual_review_continue', comments='INTERNAL REVIEW NOTE',
            override_applied=True, override_reason='INTERNAL OVERRIDE REASON'))
        db.session.commit()

    headers = auth(client)
    preview = client.get('/api/reports', headers=headers)
    assert preview.status_code == 200, preview.json
    sections = preview.json['sections']
    for internal in ('Reviews', 'Rules', 'Python predictions', 'GTM predictions'):
        assert internal not in sections
    claim_row = sections['Claims']['items'][0]
    assert 'assigned_employee_id' not in claim_row
    assert 'assigned_reviewer_id' not in claim_row
    assert 'manual_review_required' not in claim_row

    _, response = exported(app, client, headers, 'csv')
    body = response.data.decode('utf-8-sig')
    assert 'INTERNAL REVIEW NOTE' not in body
    assert 'INTERNAL OVERRIDE REASON' not in body
    assert 'Python predictions' not in body
    assert 'GTM predictions' not in body
    response.close()

    admin_preview = client.get('/api/reports', headers=auth(client, 'admin')).json
    assert admin_preview['sections']['Reviews']['items'][0]['comments'] == 'INTERNAL REVIEW NOTE'
    assert 'Rules' in admin_preview['sections']
    assert 'Python predictions' in admin_preview['sections']
    assert 'GTM predictions' in admin_preview['sections']


def test_reviewer_scope_is_assignment_only(app, client, accounts, claims):
    headers = auth(client, 'reviewer')
    claim_id = claims['customer']['claim']
    assert client.get(f'/api/reports/claim/{claim_id}', headers=headers).status_code == 404
    with app.app_context():
        claim = db.session.get(Claim, claim_id)
        claim.assigned_reviewer_id = accounts['reviewer']
        claim.status = 'manual_review'
        db.session.commit()
    data = client.get(f'/api/reports/reviewer/{accounts["reviewer"]}', headers=headers).json
    assert data['sections']['Claims']['total'] == 1
    assert client.get(f'/api/reports/customer/{accounts["customer"]}', headers=headers).status_code == 404
    assert client.get(f'/api/reports/claim/{claims["other"]["claim"]}', headers=headers).status_code == 404
    exported(app, client, headers, filters={'dataset': 'reviewer_queue'})


@pytest.mark.parametrize('format', ['csv', 'excel', 'pdf'])
def test_artifact_content_audit_and_history(app, client, accounts, claims, format):
    with app.app_context():
        claim = db.session.get(Claim, claims['customer']['claim'])
        claim.fault_description = '=HYPERLINK("https://invalid.example", "test")'
        claim.submission_date = date.today()
        db.session.add(Review(claim_id=claim.id, reviewer_user_id=accounts['reviewer'],
            decision='request_information', comments='Please provide evidence.'))
        reference = claim.claim_id
        db.session.commit()
    headers = auth(client, 'admin')
    job_id, result = exported(app, client, headers, format)
    assert 'attachment;' in result.headers['Content-Disposition']
    assert result.headers['Cache-Control'] == 'no-store'
    if format == 'csv':
        rows = list(csv.DictReader(StringIO(result.data.decode('utf-8-sig'))))
        claim = next(row for row in rows if row['section'] == 'Claims')
        assert claim['claim_id'] == reference
        assert claim['fault_description'].startswith("'=HYPERLINK")
        assert any(row['comments'] == 'Please provide evidence.' for row in rows)
    elif format == 'excel':
        workbook = load_workbook(BytesIO(result.data))
        assert {'Claims', 'Products', 'Warranties', 'Reviews', 'Analytics'} <= set(workbook.sheetnames)
        sheet = workbook['Claims']
        assert sheet.freeze_panes == 'A2' and sheet.auto_filter.ref
        assert len(sheet.conditional_formatting) > 0
        values = {sheet.cell(1, i).value: sheet.cell(2, i) for i in range(1, sheet.max_column + 1)}
        assert values['Claim Id'].value == reference
        assert values['Fault Description'].data_type == 's'
        assert values['Submission Date'].is_date
        assert workbook['Reviews']['F2'].value == 'Please provide evidence.'
        workbook.close()
    else:
        pdf = PdfReader(BytesIO(result.data))
        text = '\n'.join(page.extract_text() for page in pdf.pages)
        assert reference in text and 'Please provide evidence.' in text
        assert 'Warranties' in text and 'Analytics' in text
    history = client.get('/api/reports/history', headers=headers).json
    assert history['total'] == 1 and history['items'][0]['id'] == job_id
    other = auth(client, 'other')
    assert client.get(f'/api/reports/jobs/{job_id}', headers=other).status_code == 404
    assert client.get(f'/api/reports/jobs/{job_id}/download', headers=other).status_code == 404
    assert client.get('/api/reports/history', headers=other).json['total'] == 0
    with app.app_context():
        actions = set(db.session.scalars(select(AuditLog.action).where(AuditLog.entity_id == job_id)))
        assert {'report.requested', 'report.started', 'report.generated', 'report.download'} <= actions


def test_download_revokes_reassigned_claim_and_role_changes(app, client, accounts, claims):
    headers = auth(client, 'employee')
    job_id, response = exported(app, client, headers)
    response.close()
    with app.app_context():
        db.session.get(Claim, claims['customer']['claim']).assigned_employee_id = None
        db.session.commit()
    assert client.get(f'/api/reports/jobs/{job_id}/download', headers=headers).status_code == 403
    headers = auth(client, 'admin')
    job_id, response = exported(app, client, headers)
    response.close()
    with app.app_context():
        db.session.get(User, accounts['admin']).role = 'customer'
        db.session.commit()
    assert client.get(f'/api/reports/jobs/{job_id}/download', headers=headers).status_code == 403


def test_filters_pagination_and_expiring_warranties(app, client, accounts, claims):
    with app.app_context():
        claim = db.session.get(Claim, claims['customer']['claim'])
        claim.status, claim.submission_date = 'approved', date.today()
        db.session.get(Warranty, claims['customer']['warranty']).expiry_date = date.today() + timedelta(days=5)
        db.session.commit()
    headers = auth(client)
    filters = {'dataset': 'approved', 'search': 'Laptop', 'date_from': date.today().isoformat()}
    data = client.get('/api/reports', query_string=filters, headers=headers).json
    assert data['sections']['Claims']['total'] == 1
    _, response = exported(app, client, headers, filters=filters)
    rows = list(csv.DictReader(StringIO(response.data.decode('utf-8-sig'))))
    assert sum(row['section'] == 'Claims' for row in rows) == 1
    assert client.get('/api/reports?dataset=rejected', headers=headers).json['analytics']['total_claims'] == 0
    assert client.get('/api/reports?search=%25', headers=headers).json['analytics']['total_claims'] == 0
    page = client.get('/api/reports?page=2&per_page=1', headers=headers).json
    assert page['sections']['Claims']['items'] == [] and page['sections']['Claims']['total'] == 1
    expiring = client.get('/api/reports?dataset=expiring_warranties&expiry_days=10', headers=headers).json
    assert expiring['sections']['Warranties']['total'] == 1
    for query in ['page=0', 'per_page=101', 'date_from=bad', 'unknown=true', 'expiry_days=366', 'customer_id=999999999999999999999999',
                  'date_from=2026-10-01&date_to=2026-01-01', 'dataset=approved&status=rejected']:
        assert client.get(f'/api/reports?{query}', headers=headers).status_code == 400
    for payload in [[], {'report_type': 'claim'}, {'format': 'csv'}, {'entity_id': '../etc'}]:
        assert client.post('/api/reports/export/csv', json=payload, headers=headers).status_code == 400


def test_queue_limits_failure_recovery_and_expiration(app, client, accounts, claims, monkeypatch):
    headers = auth(client)
    app.config['REPORT_MAX_PENDING'] = 1
    job = client.post('/api/reports/export/csv', json={}, headers=headers).json
    assert client.get(f'/api/reports/jobs/{job["id"]}/download', headers=headers).status_code == 409
    assert client.post('/api/reports/export/csv', json={}, headers=headers).status_code == 429
    with app.app_context():
        from backend.services import report_jobs
        monkeypatch.setattr(report_jobs, 'write_csv', lambda *args: (_ for _ in ()).throw(RuntimeError('private path')))
        assert process_next()
        saved = db.session.get(ReportJob, job['id'])
        assert saved.status == 'failed' and 'private' not in saved.error
        saved.status = 'running'
        saved.heartbeat_at = utcnow() - timedelta(hours=1)
        db.session.commit()
        cleanup()
        assert db.session.get(ReportJob, saved.id).status == 'failed'
        saved.expires_at = utcnow() - timedelta(seconds=1)
        db.session.commit()
        cleanup()
        assert db.session.get(ReportJob, saved.id).status == 'expired'
        assert not process_next()
    assert client.get(f'/api/reports/jobs/{job["id"]}/download', headers=headers).status_code == 410


def test_inactive_requester_is_rechecked_by_worker(app, client, accounts, claims):
    headers = auth(client)
    job = client.post('/api/reports/export/excel', json={}, headers=headers).json
    with app.app_context():
        db.session.get(User, accounts['customer']).is_active = False
        db.session.commit()
        assert process_next()
        assert db.session.get(ReportJob, job['id']).status == 'failed'


def test_expiry_removes_artifact_and_abandoned_spool(app, client, accounts, claims):
    headers = auth(client)
    job_id, response = exported(app, client, headers)
    response.close()
    from backend.services.report_jobs import root, artifact
    with app.app_context():
        job = db.session.get(ReportJob, job_id)
        spool = root() / f'{job_id}-orphan'
        spool.mkdir();(spool / '0.jsonl').write_text('sensitive data', encoding='utf-8')
        path = artifact(job)
        job.expires_at = utcnow() - timedelta(seconds=1)
        db.session.commit()
        cleanup()
        assert not path.exists() and not spool.exists()
        assert db.session.scalar(select(func.count()).select_from(ReportItem).where(ReportItem.job_id == job_id)) == 0
    assert client.get(f'/api/reports/jobs/{job_id}/download', headers=headers).status_code == 410


def test_empty_reports_and_long_excel_text(app, client, accounts, claims):
    headers = auth(client)
    for format in ['csv', 'excel', 'pdf']:
        _, response = exported(app, client, headers, format, {'status': 'closed'})
        assert len(response.data) > 0
        response.close()
    with app.app_context():
        db.session.get(Claim, claims['customer']['claim']).fault_description = 'x' * 40000
        db.session.commit()
    _, response = exported(app, client, headers, 'excel')
    workbook = load_workbook(BytesIO(response.data))
    values = {workbook['Claims'].cell(1, i).value: workbook['Claims'].cell(2, i).value
              for i in range(1, workbook['Claims'].max_column + 1)}
    assert values['Fault Description'] + values['Fault Description Continued 2'] == 'x' * 40000
    workbook.close()


@pytest.mark.parametrize('text', ['=1+1', '+cmd', '-cmd', '@SUM(1)', '\t=1', ' \n=1'])
def test_spreadsheet_formulas_are_neutralized(text):
    from backend.services.report_exports import cell
    assert cell(text).startswith("'")


@pytest.mark.parametrize('format', ['csv', 'excel'])
def test_export_over_100000_records_in_bounded_batches(app, client, accounts, format):
    total = 100001
    with app.app_context():
        for first in range(0, total, 1000):
            db.session.execute(Claim.__table__.insert(), [
                {'claim_id': f'CLM-LARGE-{i}', 'user_id': accounts['customer'], 'status': 'approved'}
                for i in range(first, min(first + 1000, total))])
        db.session.commit()
    headers = auth(client)
    job_id, response = exported(app, client, headers, format, {'status': 'approved'})
    if format == 'csv':
        rows = csv.DictReader(StringIO(response.data.decode('utf-8-sig')))
        assert sum(row['section'] == 'Claims' for row in rows) == total
    else:
        book = load_workbook(BytesIO(response.data), read_only=True)
        assert sum(1 for _ in book['Claims'].rows) == total + 1
        book.close()
    with app.app_context():
        assert db.session.scalar(select(func.count()).select_from(ReportItem).where(
            ReportItem.job_id == job_id, ReportItem.kind == 'claim')) == total
        job = db.session.get(ReportJob, job_id)
        assert job.processed >= total


def test_report_migration_roundtrip(app):
    from alembic import command
    from alembic.config import Config
    from sqlalchemy import inspect
    config = Config('backend/db/alembic.ini')
    with app.app_context():
        assert 'report_jobs' in inspect(db.engine).get_table_names()
        db.session.remove()
        command.downgrade(config, 'a194c5e72d31')
        assert 'report_jobs' not in inspect(db.engine).get_table_names()
        command.upgrade(config, 'head')
        assert 'report_items' in inspect(db.engine).get_table_names()
