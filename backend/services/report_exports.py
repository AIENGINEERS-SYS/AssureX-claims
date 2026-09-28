"""Disk-spooled writers with bounded row buffers and write-only Excel worksheets."""
import csv
from datetime import date, datetime
import json
import re

CONTROL = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def cell(value):
    if value is None:
        return ''
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    if isinstance(value, str):
        value = CONTROL.sub('', value)
        if value.lstrip().startswith(('=', '+', '-', '@', '\t', '\r', '\n')) or value.startswith(('\t', '\r', '\n')):
            value = "'" + value
    return value


def rows(spec):
    with spec['path'].open(encoding='utf-8') as source:
        for line in source:
            yield json.loads(line)


def write_csv(path, specs, heartbeat):
    columns = list(dict.fromkeys(key for spec in specs.values() for key in spec['columns']))
    with path.open('w', encoding='utf-8-sig', newline='') as output:
        writer = csv.writer(output)
        writer.writerow(['section', *columns])
        for name, spec in specs.items():
            for i, row in enumerate(rows(spec)):
                writer.writerow([name, *(cell(row.get(key)) for key in columns)])
                if i % 1000 == 0:
                    heartbeat()


def write_excel(path, specs, heartbeat):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment
    workbook = Workbook(write_only=True)
    alignment = Alignment(vertical='top', wrap_text=True)
    try:
        for name, spec in specs.items():
            sheet = None
            for i, row in enumerate(rows(spec)):
                # Excel has 1,048,576 rows including the header. Split without truncation.
                if i % 1048575 == 0:
                    sheet = _sheet(workbook, name, spec, i // 1048575)
                output = []
                for key in spec['columns']:
                    value = cell(row.get(key))
                    if key in spec['dates'] and value:
                        value = datetime.fromisoformat(value).replace(tzinfo=None) if 'T' in value else date.fromisoformat(value)
                    # Keep overlong text in continuation columns rather than silently losing it.
                    values = [value]
                    if key in spec.get('long_columns', {}):
                        text = str(value)
                        values = [text[n:n + 32000] for n in range(0, len(text), 32000)] or ['']
                        values += [''] * (spec['long_columns'][key] - len(values))
                    for part in values:
                        item = WriteOnlyCell(sheet, value=part)
                        if isinstance(part, str):
                            item.data_type = 's'
                        if isinstance(part, (date, datetime)):
                            item.number_format = 'yyyy-mm-dd hh:mm:ss' if isinstance(part, datetime) else 'yyyy-mm-dd'
                        item.alignment = alignment
                        output.append(item)
                sheet.append(output)
                if i % 1000 == 0:
                    heartbeat()
            if sheet is None:
                _sheet(workbook, name, spec, 0)
        workbook.save(path)
    finally:
        # openpyxl owns temporary worksheet files; close and remove them on failure too.
        for sheet in workbook.worksheets:
            if not sheet.closed:
                sheet.close()
            writer = getattr(sheet, '_writer', None)
            if writer and __import__('os').path.exists(writer.out):
                writer.cleanup()
        workbook.close()


def _sheet(workbook, name, spec, number):
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.utils import get_column_letter
    sheet = workbook.create_sheet(name if not number else f'{name[:24]} {number + 1}')
    sheet.freeze_panes = 'A2'
    headers = []
    for key in spec['columns']:
        for part in range(spec.get('long_columns', {}).get(key, 1)):
            headers.append((key, key if part == 0 else f'{key} continued {part + 1}'))
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{min(spec["count"] + 1, 1048576)}'
    for i, (key, title) in enumerate(headers, 1):
        letter = get_column_letter(i)
        sheet.column_dimensions[letter].width = min(65, max(len(title) + 2, spec['widths'].get(key, 12) + 2))
        if key in {'status', 'decision', 'result'}:
            for status, color in [('approved', 'DCFCE7'), ('approve', 'DCFCE7'), ('passed', 'DCFCE7'),
                                  ('rejected', 'FEE2E2'), ('reject', 'FEE2E2'), ('failed', 'FEE2E2'),
                                  ('manual_review', 'FEF3C7')]:
                sheet.conditional_formatting.add(f'{letter}2:{letter}1048576',
                    CellIsRule(operator='equal', formula=[f'"{status}"'], fill=PatternFill('solid', fgColor=color)))
    output = []
    for _, title in headers:
        item = WriteOnlyCell(sheet, value=title.replace('_', ' ').title())
        item.font = Font(bold=True, color='FFFFFF')
        item.fill = PatternFill('solid', fgColor='153C35')
        output.append(item)
    sheet.append(output)
    return sheet


def write_pdf(path, specs, heartbeat):
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from flask import current_app
    from pathlib import Path
    import reportlab
    font = current_app.config.get('REPORT_PDF_FONT') or str(Path(reportlab.__file__).parent / 'fonts' / 'Vera.ttf')
    import hashlib
    face = 'AssureXReport' + hashlib.sha256(str(font).encode()).hexdigest()[:12]
    if face not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(face, font))
    pdf = canvas.Canvas(str(path), pagesize=A4, pageCompression=1)
    pdf.setTitle('AssureX report')
    y, page = 0, 0

    def line(text, heading=False):
        nonlocal y, page
        if y < 55:
            if page:
                pdf.showPage()
            page += 1
            pdf.setFont(face, 9)
            pdf.drawString(40, 810, 'AssureX | Confidential report')
            pdf.drawRightString(555, 30, f'Page {page}')
            y = 780
        pdf.setFont(face, 11 if heading else 9)
        pdf.drawString(40, y, text)
        y -= 16 if heading else 12

    for name, spec in specs.items():
        line(f'{name} ({spec["count"]} records)', True)
        for i, row in enumerate(rows(spec)):
            line(f'Record {i + 1}', True)
            for key in spec['columns']:
                value = row.get(key)
                if value is None:
                    continue
                if isinstance(value, (dict, list)):
                    value = json.dumps(value, ensure_ascii=False)
                text = CONTROL.sub('', str(value))
                for paragraph in (f'{key.replace("_", " ")}: {text}').splitlines():
                    # Width measured using the selected font, so long Unicode values cannot overrun.
                    remaining = paragraph
                    while remaining:
                        end = min(len(remaining), 110)
                        while end > 1 and pdfmetrics.stringWidth(remaining[:end], face, 9) > 510:
                            end -= 1
                        line(remaining[:end])
                        remaining = remaining[end:]
            if i % 250 == 0:
                heartbeat()
    if not page:
        line('No records matched this report.')
    pdf.save()
